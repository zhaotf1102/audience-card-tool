from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

CATEGORIES = [
    "男高",
    "男富",
    "男帅",
    "女白",
    "女富",
    "女美",
]


def main() -> None:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for title in CATEGORIES:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.25), Inches(12), Inches(0.8))
        p = box.text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(34)
        p.font.bold = True
    out = Path("templates/demo_template.pptx")
    out.parent.mkdir(exist_ok=True)
    prs.save(out)
    print(f"created {out}")


if __name__ == "__main__":
    main()
