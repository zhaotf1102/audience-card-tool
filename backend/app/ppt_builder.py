from __future__ import annotations

import copy
from pathlib import Path
from typing import Any, Dict, List

from pptx import Presentation
from pptx.util import Inches

# MVP fixed template slide mapping. This is 1-based for humans.
# User can modify later or expose in UI.
DEFAULT_TEMPLATE_SLIDE_INDEX = {
    "male_tall": 1,
    "male_rich": 2,
    "male_handsome": 3,
    "female_fair": 4,
    "female_rich": 5,
    "female_beautiful": 6,
}


def build_ppt_from_crops(template_path: Path, output_path: Path, state: Dict[str, Any], crops_dir: Path) -> Path:
    """Build a PPTX with one cropped card image per slide.

    The MVP uses fixed template slide indexes for each category. Each generated slide
    is a duplicate of that category's template slide, then one crop image is placed
    in a large centered safe area.

    TODO: Replace fixed mapping with template notes / hidden placeholders.
    """
    prs = Presentation(str(template_path))

    category_to_crops: Dict[str, List[Path]] = {}
    for card in state.get("cards", []):
        category = card.get("category")
        if not category:
            continue
        for url in card.get("crop_urls", []):
            filename = url.rstrip("/").split("/")[-1]
            path = crops_dir / filename
            if path.exists():
                category_to_crops.setdefault(category, []).append(path)

    # Remove original category template slides? For MVP, keep original template slides
    # and append generated pages after them. This avoids breaking complex templates.
    for category, crop_paths in category_to_crops.items():
        template_number = DEFAULT_TEMPLATE_SLIDE_INDEX.get(category)
        if not template_number:
            continue
        template_idx = template_number - 1
        if template_idx >= len(prs.slides):
            continue
        template_slide = prs.slides[template_idx]
        for crop_path in crop_paths:
            new_slide = duplicate_slide(prs, template_slide)
            add_centered_image(prs, new_slide, crop_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path


def duplicate_slide(prs: Presentation, source_slide):
    blank_layout = prs.slide_layouts[6]
    dest = prs.slides.add_slide(blank_layout)

    for shape in source_slide.shapes:
        new_element = copy.deepcopy(shape.element)
        dest.shapes._spTree.insert_element_before(new_element, "p:extLst")

    # Copy relationships such as images/backgrounds where python-pptx permits.
    # This minimal approach is usually enough for shape/text-based templates.
    for rel in source_slide.part.rels.values():
        if "notesSlide" in rel.reltype:
            continue
        try:
            dest.part.rels.add_relationship(rel.reltype, rel._target, rel.rId)
        except Exception:
            pass

    return dest


def add_centered_image(prs: Presentation, slide, image_path: Path) -> None:
    slide_w = prs.slide_width
    slide_h = prs.slide_height

    # Large safe area. Leave title and margins visible.
    left = Inches(0.7)
    top = Inches(1.25)
    max_w = slide_w - Inches(1.4)
    max_h = slide_h - Inches(1.65)

    pic = slide.shapes.add_picture(str(image_path), left, top)
    ratio = min(max_w / pic.width, max_h / pic.height)
    pic.width = int(pic.width * ratio)
    pic.height = int(pic.height * ratio)
    pic.left = int((slide_w - pic.width) / 2)
    pic.top = int(top + (max_h - pic.height) / 2)
